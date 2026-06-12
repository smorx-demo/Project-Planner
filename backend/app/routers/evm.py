from fastapi import APIRouter, Depends, HTTPException, Request, Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from pydantic import BaseModel
from typing import Optional
import uuid
from datetime import date, timedelta, datetime as dt
from app.database import get_db
from app.models.activity import Activity
from app.models.project import Project
from app.models.evm_snapshot import EVMSnapshot
from app.models.user import User
from app.dependencies import get_current_user, require_manager, check_project_access
from app.services.evm_service import EVMService, EPOCH as EVM_EPOCH
from app.services.status_service import StatusService
from app.envelope import ok, audit
from app.limiter import limiter
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO

router = APIRouter(tags=["evm"])


class EVMUpdatePayload(BaseModel):
    bac:         Optional[float] = None
    bac_unit:    Optional[str]   = None
    actual_pct:  Optional[float] = None
    actual_cost: Optional[float] = None
    planned_pct: Optional[float] = None


@router.get("/projects/{project_id}/evm")
@limiter.limit("60/minute")
async def get_project_evm(
    request: Request,
    project_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    await check_project_access(project_id, current_user, db)
    result = await db.execute(
        select(Activity)
        .where(Activity.project_id == project_id)
        .order_by(Activity.sequence_no)
    )
    activities = result.scalars().all()
    return ok(EVMService.compute_project_evm(activities))


@router.get("/projects/{project_id}/evm/history")
@limiter.limit("60/minute")
async def get_evm_history(
    request: Request,
    project_id: uuid.UUID,
    days: int = 90,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    await check_project_access(project_id, current_user, db)
    since = date.today() - timedelta(days=days)
    result = await db.execute(
        select(EVMSnapshot)
        .where(
            EVMSnapshot.project_id == project_id,
            EVMSnapshot.activity_id == None,   # noqa: E711
            EVMSnapshot.snapshot_date >= since,
        )
        .order_by(EVMSnapshot.snapshot_date.asc())
    )
    snaps = result.scalars().all()
    return ok([
        {
            "snapshot_date": str(s.snapshot_date),
            "pv": s.pv, "ev": s.ev, "ac": s.ac,
            "spi": s.spi, "cpi": s.cpi, "sv": s.sv, "cv": s.cv,
        }
        for s in snaps
    ])


@router.get("/activities/{activity_id}/evm")
@limiter.limit("60/minute")
async def get_activity_evm(
    request: Request,
    activity_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    activity = await db.get(Activity, activity_id)
    if not activity:
        raise HTTPException(status_code=404, detail="Activity not found")
    return ok(EVMService.compute_activity_evm(activity))


@router.put("/activities/{activity_id}/evm")
@limiter.limit("60/minute")
async def update_activity_evm(
    request: Request,
    activity_id: uuid.UUID,
    payload: EVMUpdatePayload,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_manager),
):
    activity = await db.get(Activity, activity_id)
    if not activity:
        raise HTTPException(status_code=404, detail="Activity not found")

    old_vals = {}
    if payload.bac is not None:
        old_vals['bac'] = activity.bac
        activity.bac = payload.bac
    if payload.bac_unit is not None:
        old_vals['bac_unit'] = activity.bac_unit
        activity.bac_unit = payload.bac_unit
    if payload.actual_pct is not None:
        old_vals['actual_pct'] = activity.actual_pct
        activity.actual_pct = max(0.0, min(100.0, payload.actual_pct))
    if payload.actual_cost is not None:
        old_vals['actual_cost'] = activity.actual_cost
        activity.actual_cost = payload.actual_cost
    if payload.planned_pct is not None:
        old_vals['planned_pct'] = activity.planned_pct
        activity.planned_pct = max(0.0, min(100.0, payload.planned_pct))

    await db.flush()
    await db.refresh(activity)

    await audit(
        db,
        user_id=current_user.id,
        action="UPDATE",
        entity_type="activity_evm",
        entity_id=str(activity_id),
        project_id=activity.project_id,
        old_value={k: str(v) for k, v in old_vals.items()},
        new_value={k: str(v) for k, v in payload.model_dump(exclude_none=True).items()},
    )
    return ok(EVMService.compute_activity_evm(activity))


@router.get("/projects/{project_id}/pms")
@limiter.limit("60/minute")
async def get_pms(
    request: Request,
    project_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    await check_project_access(project_id, current_user, db)

    result = await db.execute(
        select(Activity)
        .where(Activity.project_id == project_id)
        .order_by(Activity.sequence_no)
    )
    activities = result.scalars().all()

    def col_to_str(col):
        if col is None:
            return None
        return str(EVM_EPOCH + timedelta(days=col))

    rows = []
    for idx, a in enumerate(activities, 1):
        evm = EVMService.compute_activity_evm(a)
        pv_pct = evm["pv_pct"] or 0
        rows.append({
            "seq":          idx,
            "name":         a.name,
            "group_type":   a.group_type,
            "plan_start":   col_to_str(a.plan_start_col) or "",
            "plan_end":     col_to_str(a.plan_end_col) or "",
            "actual_start": col_to_str(a.actual_start_col),
            "actual_end":   col_to_str(a.actual_end_col),
            "bac":          a.bac,
            "bac_unit":     a.bac_unit or "",
            "planned_pct":  round(pv_pct / 100, 4),
            "actual_pct":   a.actual_pct,
            "actual_cost":  a.actual_cost,
            "pv":           evm["pv"],
            "ev":           evm["ev"],
            "ac":           evm["ac"],
            "spi":          evm["spi"],
            "cpi":          evm["cpi"],
            "status":       StatusService.compute_status(a),
        })

    return ok(rows)


@router.get("/projects/{project_id}/pms/export")
@limiter.limit("20/minute")
async def export_pms(
    request: Request,
    project_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    await check_project_access(project_id, current_user, db)

    result = await db.execute(
        select(Activity)
        .where(Activity.project_id == project_id)
        .order_by(Activity.sequence_no)
    )
    activities = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Progress Measurement Sheet"

    # Color fills
    navy_fill   = PatternFill("solid", fgColor="1F4E79")
    spi_warn    = PatternFill("solid", fgColor="FFF2CC")
    cpi_warn    = PatternFill("solid", fgColor="FCE4D6")

    # Header block (rows 1-5)
    meta = [
        ("Project",     getattr(project, "name", "") or ""),
        ("PO Number",   getattr(project, "po_number", "") or ""),
        ("Part Number", getattr(project, "part_number", "") or ""),
        ("Customer",    getattr(project, "customer_name", "") or ""),
        ("Generated",   dt.now().strftime("%d %b %Y %H:%M")),
    ]
    for i, (k, v) in enumerate(meta, 1):
        ws.cell(i, 1, k).font = Font(bold=True, size=10)
        ws.cell(i, 2, v)
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 36

    # Table headers (row 7)
    COLS = ["Sn", "Activity", "Group", "BAC", "Unit", "Planned %",
            "Actual %", "Variance %", "PV", "EV", "AC", "SPI", "CPI", "Status"]
    HDR_ROW = 7
    for ci, h in enumerate(COLS, 1):
        cell = ws.cell(HDR_ROW, ci, h)
        cell.fill = navy_fill
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[HDR_ROW].height = 20

    # Column widths
    col_widths = [5, 40, 16, 10, 8, 12, 12, 12, 10, 10, 10, 8, 8, 12]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # Status fill map
    STATUS_FILL = {
        "DONE":     PatternFill("solid", fgColor="E2EFDA"),
        "DELAYED":  PatternFill("solid", fgColor="FCE4D6"),
        "SLOW":     PatternFill("solid", fgColor="FFF2CC"),
        "ON_TRACK": PatternFill("solid", fgColor="DBEAFE"),
        "PENDING":  PatternFill("solid", fgColor="F1F5F9"),
    }

    for ri, a in enumerate(activities, HDR_ROW + 1):
        evm = EVMService.compute_activity_evm(a)
        status = StatusService.compute_status(a)
        row_fill = STATUS_FILL.get(status)

        values = [
            a.sequence_no, a.name, a.group_type or "",
            a.bac or "", a.bac_unit,
            round(evm["pv_pct"] or 0, 1),
            round(a.actual_pct or 0, 1),
            round(((a.actual_pct or 0) - (evm["pv_pct"] or 0)), 1),
            evm["pv"] or "", evm["ev"] or "", evm["ac"] or "",
            evm["spi"] or "", evm["cpi"] or "",
            status,
        ]
        for ci, val in enumerate(values, 1):
            cell = ws.cell(ri, ci, val)
            if row_fill:
                cell.fill = row_fill
            cell.alignment = Alignment(
                horizontal="center" if ci not in (2, 3) else "left"
            )
            # SPI < 1 warning (col 12)
            if ci == 12 and isinstance(val, float) and val < 1:
                cell.fill = spi_warn
            # CPI < 1 warning (col 13)
            if ci == 13 and isinstance(val, float) and val < 1:
                cell.fill = cpi_warn

    # Footer TOTAL row
    footer_row = HDR_ROW + len(activities) + 1
    project_evm = EVMService.compute_project_evm(activities)
    footer_vals = [
        "", "TOTAL", "", project_evm["total_bac"], "",
        "", "", "",
        project_evm["total_pv"], project_evm["total_ev"], project_evm["total_ac"],
        project_evm["spi"] or "", project_evm["cpi"] or "", "",
    ]
    for ci, val in enumerate(footer_vals, 1):
        cell = ws.cell(footer_row, ci, val)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(
            horizontal="center" if ci not in (2, 3) else "left"
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"PMS_{getattr(project, 'po_number', None) or project_id}_{dt.now().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
